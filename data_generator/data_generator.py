#!/usr/bin/env python
import os
import tempfile
import shutil
import requests
from openpyxl import load_workbook
from transliterate import translit
import argparse
from jinja2 import Environment, FileSystemLoader


SHEET_URL = 'https://docs.google.com/spreadsheets/d/1HQLdWm3ZSBvlqHKekofgxMuhDABR4HL_TFQ4AyGM5xA/export'

template_env = Environment(
    loader=FileSystemLoader('templates'),
    autoescape=False
)


def render_template(template_name, context, output_file):
    template = template_env.get_template(template_name)
    rendered = template.render(context)
    with open(output_file, "w") as fh:
        fh.write(rendered)


def mkdir(d):
    try:
        os.makedirs(d)
    except FileExistsError:
        pass


def download_sheet():
    (fp, filename) = tempfile.mkstemp(suffix='.xlsx')
    r = requests.get(SHEET_URL, stream=True)
    r.raise_for_status()
    r.raw.decode_content = True
    with open(fp, 'wb', closefd=True) as f:
        shutil.copyfileobj(r.raw, f)
    return filename


def read_sheet():
    filename = download_sheet()
    wb = load_workbook(filename=filename, read_only=True)
    load_words = True
    words = {}
    lessons = {}
    for sheet_name in wb.get_sheet_names():
        ws = wb[sheet_name]

        if load_words:
            for row in ws.iter_rows(min_row=2):
                (word, word_type, pronounce_as, image) = [row[idx].value if idx < len(row) else None for idx in range(4)]
                if not word:
                    break
                if not word_type or not image:
                    print(f"type or image not defined for '{word}', skipping!")
                    continue
                words[word] = {
                    'word': word,
                    'type': word_type,
                    'pronounce_as': pronounce_as,
                    'image': image,
                }
            load_words = False
        else:
            lessons[sheet_name] = {
                'filename': lesson_filename(sheet_name),
                'columns': [[], [], [], []],

            }
            for row in ws.iter_rows(min_row=2):
                columns = [row[idx].value if idx < len(row) else None for idx in range(4)]
                if columns == [None, None, None, None]:
                    break
                for idx in range(len(columns)):
                    if columns[idx]:
                        lessons[sheet_name]['columns'][idx].append(columns[idx])
    os.unlink(filename)
    return (words, lessons)


def lesson_filename(lesson_name):
    lesson_name = translit(lesson_name, 'bg', reversed=True)
    lesson_name = lesson_name.lower().replace(' ', '-')
    return f'{lesson_name}.js'


def convert_images(words, img_dir, output_dir):
    img_out_dir = os.path.join(output_dir, 'img')
    mkdir(img_out_dir)

    for word in words.values():
        image_filename = word['image']
        filename, file_extension = os.path.splitext(image_filename)
        in_file = os.path.join(img_dir, image_filename)
        if not os.path.exists(in_file):
            raise Exception(f'File missing: {in_file}')
        out_file = f'{filename}.png'
        if file_extension == '.svg':
            cmd = f'rsvg-convert --background-color=none -o {img_out_dir}/{out_file} -w 50 -h 50 {in_file}'
        else:
            cmd = f'gm convert {in_file} -resize 50x50 {img_out_dir}/{out_file}'
        print('Run:', cmd)
        os.system(cmd)
        word['image_png'] = f'img/{out_file}'


def write_lesson_index(lessons, output_dir):
    context = {
        'lessons': lessons
    }
    render_template('lesson_index.jinja2', context, os.path.join(output_dir, 'lesson-index.js'))


def write_lesson(lesson, words, output_dir):
    columns = []
    for column_words in lesson['columns']:
        if not column_words:
            continue
        column = []
        for word in column_words:
            column.append(words[word])
        columns.append(column)
    context = {
        "columns": columns
    }
    render_template('lesson_template.jinja2', context, os.path.join(output_dir, lesson['filename']))


def write_words_categories(words, output_dir):
    categories = {}
    for word in words.values():
        if word['type'] not in categories:
            categories[word['type']] = {}
        categories[word['type']][word['word']] = word
    context = {
        'categories': categories
    }
    render_template('words_categories.jinja2', context, os.path.join(output_dir, 'words-categories.js'))


parser = argparse.ArgumentParser(description='Import google docs sheet to lessons')
parser.add_argument('output_dir', help='output dir')
parser.add_argument('src_images', help='source images dir')
parser.add_argument('--rm', dest='remove', action='store_true', help='delete output dir before generating')


args = parser.parse_args()

words, lessons = read_sheet()
# print('WORDS')
# print(words)
# print('LESSONS')
# print(lessons)
if args.remove:
    print('Deleting dir:', args.output_dir)
    shutil.rmtree(args.output_dir)

print('Writing data in:', args.output_dir)
mkdir(args.output_dir)

convert_images(words, args.src_images, args.output_dir)
for lesson_name, lesson in lessons.items():
    write_lesson(lesson, words, args.output_dir)
write_lesson_index(lessons, args.output_dir)
write_words_categories(words, args.output_dir)
